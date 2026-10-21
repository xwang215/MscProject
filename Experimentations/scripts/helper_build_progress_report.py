from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[1]
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.65)
sec.left_margin = sec.right_margin = Inches(.7)
for name in ['Normal', 'Title', 'Heading 1']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
normal = doc.styles['Normal']
normal.font.size = Pt(10)
normal.paragraph_format.line_spacing = 1.05
normal.paragraph_format.space_after = Pt(4)
doc.styles['Title'].font.size = Pt(17)
doc.styles['Title'].paragraph_format.space_after = Pt(3)
doc.styles['Heading 1'].font.size = Pt(11)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(7)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(4)
md = []
def pair(en, zh, style=None):
    p = doc.add_paragraph(en, style)
    p.paragraph_format.keep_with_next = True
    doc.add_paragraph(zh, style)
    md.extend([en, zh])
def heading(en, zh):
    p = doc.add_paragraph(en + '  ' + zh, 'Heading 1')
    md.append('## ' + en + ' / ' + zh)

pair('Ulcerative Colitis Severity Estimation', '溃疡性结肠炎严重程度估计', 'Title')
pair('Research progress report | 21 October 2026', '研究进展报告｜2026 年 10 月 21 日')
heading('Objective and completed work', '研究目标与已完成工作')
pair('We established a reproducible baseline comparing a CNN with two vision Transformers for image-based UC severity classification. The completed work covers data preparation, label auditing, model evaluation, and bilingual documentation. It provides an experimental starting point, not clinical validation.',
'我们建立了可复现的基线，比较一个 CNN 与两个视觉 Transformer 对 UC 图像严重程度的分类表现。已完成数据准备、标签审计、模型评估及双语文档，为后续实验提供基础，目前尚未开展临床验证。')
heading('Dataset and preparation', '数据集与预处理')
pair('The labelled HyperKvasir archive was downloaded and verified. Its UC subset contains 851 images across six labels. The primary experiment uses 777 images with definite grades 1, 2, and 3; 74 interval-labelled images are excluded without relabelling. The training, validation, and test sets contain 465, 156, and 156 images.',
'已下载并校验带标签的 HyperKvasir 数据包。UC 子集包含六种标签、共 851 张图像。主要实验使用明确标注为 1、2、3 级的 777 张图像；排除 74 张区间标签图像，未重新标注。训练、验证和测试集分别包含 465、156 和 156 张图像。')
pair('Image readability and duplicate-label consistency were checked. Perceptually similar images were grouped within the same split to reduce leakage risk. Patient identifiers are unavailable, so patient-independent evaluation cannot be confirmed. RGB conversion, resizing, cropping, and normalization follow each model’s input requirements.',
'已检查图像可读性及重复图像标签的一致性，并将感知相似图像保留在同一数据子集中，以降低泄漏风险。由于缺少患者标识，无法确认患者层面的独立评估。RGB 转换、缩放、裁剪及归一化遵循各模型的输入要求。')
heading('Methods and observed results', '方法与实测结果')
pair('ImageNet-pretrained ResNet-18, DeiT-Tiny, and Swin-Tiny were frozen. Their features were standardized and used to train class-balanced logistic regression classifiers on CPU. Hyperparameters were selected using validation macro-F1; all models used the same test images. No end-to-end fine-tuning was performed.',
'冻结 ImageNet 预训练的 ResNet-18、DeiT-Tiny 和 Swin-Tiny，标准化其特征，并在 CPU 上训练带类别平衡权重的逻辑回归分类器。使用验证集宏平均 F1 选择超参数，所有模型使用相同测试图像，尚未进行端到端微调。')
table = doc.add_table(rows=1, cols=4)
table.style = 'Table Grid'
for cell, text in zip(table.rows[0].cells, ['Model\n模型', 'Accuracy\n准确率', 'Macro F1\n宏平均 F1', 'Grade 3 recall\n3 级召回率']):
    cell.text = text
    sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'E8EEF3'); cell._tc.get_or_add_tcPr().append(sh)
for row in [('ResNet-18','67.31%','0.6620','74.07%'),('DeiT-Tiny','66.03%','0.6617','81.48%'),('Swin-Tiny','73.08%','0.7196','66.67%')]:
    for c, text in zip(table.add_row().cells, row): c.text = text
for row in table.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.space_before = Pt(3)
            for r in p.runs: r.font.size = Pt(9)
pair('Swin-Tiny achieved the highest observed macro-F1, but DeiT-Tiny had the highest grade-3 recall. Aggregate performance therefore did not identify the most sensitive model for severe cases. These results do not establish statistical superiority or a general Transformer advantage.',
'Swin-Tiny 的实测宏平均 F1 最高，但 DeiT-Tiny 的 3 级召回率最高。因此，综合表现最佳不代表对重度病例最敏感。这些结果尚不能证明统计学优势，也不能证明 Transformer 普遍优于 CNN。')
doc.add_page_break()
heading('Labeling and research contribution', '标注工作与研究贡献')
pair('Labeling-related work reused existing image-level grades, audited their distribution and consistency, and preserved image–label mappings in split manifests. No new expert labels, bounding boxes, or segmentation masks were created. Model predictions were evaluated against reference labels; they were not treated as verified annotations.',
'标注相关工作包括复用已有图像级等级、检查标签分布与一致性，以及在数据划分清单中保留图像与标签的对应关系。未创建新的专家标签、边界框或分割掩膜。模型预测仅用于与参考标签比较，未被视为经过验证的标注。')
pair('The current contribution is a traceable comparison with saved code, data manifests, predictions, and bilingual reports. It also exposes a trade-off between overall classification performance and severe-grade sensitivity. Reduced annotation cost and improved clinical performance have not yet been demonstrated.',
'当前贡献是建立可追溯的比较流程，保存代码、数据清单、预测结果及双语报告，并揭示整体分类表现与重度等级敏感性之间的权衡。尚未证明能够降低标注成本或改善临床表现。')
heading('Relationship to prior research', '与既有研究的关系')
pair('The review of 12 studies places this work closest to three-grade CNN classification and CNN–Transformer comparisons. UC-NfNet already compared CNNs and Transformers, while UC-former evaluated Transformer-based grading with external validation [1, 2]. Thus, using Transformers alone is not a novel contribution. Ordinal loss design offers a relevant next comparator [3].',
'对 12 篇研究的比较表明，本项目最接近三级 CNN 分类及 CNN 与 Transformer 比较。UC-NfNet 已比较这两类模型，UC-former 已开展 Transformer 分级及外部验证 [1, 2]。因此，仅使用 Transformer 不构成创新。有序损失函数可作为下一阶段的重要对照 [3]。')
heading('Limitations and next steps', '局限性与下一步计划')
pair('The study uses a small, imbalanced dataset without definite grade 0 or external validation. Interval labels are excluded, and visible overlays may confound predictions. Model size and preprocessing differ. Frozen-feature results cannot establish the performance of fully fine-tuned networks or the full clinical severity of UC.',
'本研究数据量较小且类别不平衡，缺少明确的 0 级及外部验证。区间标签被排除，图像叠加信息可能干扰预测；模型规模与预处理也存在差异。冻结特征实验不能代表充分微调网络的表现，也不能代表 UC 的完整临床严重程度。')
pair('Next, establish matched end-to-end supervised baselines and test self-supervised learning (SSL) under fixed label budgets. Compare ordinary and ordinal objectives. If expert UC lesion masks become available, assess few-shot segmentation (FSS) and its added value for grading, counting mask-annotation costs separately. SSL, FSS, YOLO-assisted labeling, and VQA remain proposed, not completed, experiments.',
'下一步先建立条件匹配的端到端监督基线，再在固定标注预算下测试自监督学习（SSL），并比较普通与有序目标函数。若获得专家 UC 病灶掩膜，再评估少样本分割（FSS）及其对分级的增益，单独统计掩膜标注成本。SSL、FSS、YOLO 辅助标注和 VQA 均为拟开展、尚未完成的实验。')
heading('Selected references', '主要参考文献')
pair('[1] Turan and Durmus (2022). UC-NfNet. doi:10.1016/j.media.2022.102587\n[2] Qi et al. (2023). UC-former study. doi:10.1177/17562848231170945\n[3] Polat et al. (2022). Class Distance Weighted Cross-Entropy Loss. arXiv:2202.05167',
'[1] Turan 和 Durmus（2022）：UC-NfNet 研究。\n[2] Qi 等人（2023）：UC-former 研究。\n[3] Polat 等人（2022）：类别距离加权交叉熵损失研究。')
doc.core_properties.title = 'Ulcerative Colitis Severity Estimation Research Progress'
doc.core_properties.author = ''
for element in [doc.styles.element, doc.element]:
    for border in list(element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
for cell in table._cells:
    borders = OxmlElement('w:tcBorders')
    for edge in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement('w:' + edge)
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:color'), 'D9D9D9')
        borders.append(border)
    cell._tc.get_or_add_tcPr().append(borders)
doc.save(root / 'reports/progress_summary_bilingual_2pages.docx')
